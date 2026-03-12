import os
import tempfile

import pytest
from pptx import Presentation
from pptx.shapes.group import GroupShape

from converter import convert
from converter.parser import DrawioParser

FIXTURE = os.path.join(os.path.dirname(__file__), "fixtures", "grouped_diagram.drawio")


# --- Parser tests ---

def test_parser_marks_group_cell():
    pages = DrawioParser(FIXTURE).parse()
    vertices, _ = pages[0]["data"]
    groups = [v for v in vertices if v["is_group"]]
    assert len(groups) == 1
    assert groups[0]["id"] == "g1"


def test_parser_non_group_cells_not_marked():
    pages = DrawioParser(FIXTURE).parse()
    vertices, _ = pages[0]["data"]
    non_groups = [v for v in vertices if not v["is_group"]]
    assert len(non_groups) == 3  # a, b, c


def test_parser_group_child_absolute_positions():
    pages = DrawioParser(FIXTURE).parse()
    vertices, _ = pages[0]["data"]
    a = next(v for v in vertices if v["id"] == "a")
    b = next(v for v in vertices if v["id"] == "b")
    assert a["x"] == 110
    assert a["y"] == 120
    assert b["x"] == 210
    assert b["y"] == 120


# --- Engine/integration tests ---

@pytest.fixture
def output_pptx(tmp_path):
    out = str(tmp_path / "out.pptx")
    convert(FIXTURE, out)
    return Presentation(out)


def test_slide_contains_one_group(output_pptx):
    slide = output_pptx.slides[0]
    groups = [s for s in slide.shapes if isinstance(s, GroupShape)]
    assert len(groups) == 1


def test_group_contains_two_child_shapes(output_pptx):
    slide = output_pptx.slides[0]
    group = next(s for s in slide.shapes if isinstance(s, GroupShape))
    assert len(group.shapes) == 2


def test_top_level_shape_exists_outside_group(output_pptx):
    slide = output_pptx.slides[0]
    non_groups = [s for s in slide.shapes if not isinstance(s, GroupShape)]
    # Should contain: shape C + connector e1 (edge labels excluded in this diagram)
    assert len(non_groups) >= 2  # at minimum C and the connector
